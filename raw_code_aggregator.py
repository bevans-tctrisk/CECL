"""Aggregate charge-offs / recoveries / balances directly from a WARM's raw
per-record loan-code tabs.

Used when loan codes are broken out into new pools: the WARM's aggregated
Display / HistBal tabs group or mislabel those pools, so each broken-out pool's
history is re-derived from the raw per-record tabs keyed by loan code:
  - ``CO Data``    : charge-off records (loan code col C, amount col D, date col E)
  - ``Recov Data`` : recovery records  (loan code col C, amount col D, date col I)
  - ``<Mon>-<YY> Data`` quarterly snapshots (loan code col C, balance col D,
    current FICO col F)
"""
from __future__ import annotations

import datetime
import re
from collections import defaultdict

import openpyxl
import pandas as pd

_MON = {'Jan': 1, 'Feb': 2, 'Mar': 3, 'Apr': 4, 'May': 5, 'Jun': 6,
        'Jul': 7, 'Aug': 8, 'Sep': 9, 'Oct': 10, 'Nov': 11, 'Dec': 12}


def _to_dt(v):
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, (int, float)) and v > 40000:
        return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=int(v))
    return None


def _year(v):
    d = _to_dt(v)
    return d.year if d else None


def _ym(v):
    d = _to_dt(v)
    return (d.year, d.month) if d else None


def _int_code(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _band(score, grades, no_score='Not Reported'):
    try:
        s = float(score)
    except (TypeError, ValueError):
        return no_score
    if s <= 0:
        return no_score
    for g in grades:
        lo, hi = g.get('min_score'), g.get('max_score')
        if lo is not None and hi is not None and lo <= s <= hi:
            return g.get('label')
    return no_score


def aggregate(warm_path, pool_codes, grades, date_grid, snapshot_date,
              df=None, acl_months=None, no_score='Not Reported'):
    """Aggregate CO/Rc/balances by loan code from ``warm_path``'s raw tabs.

    ``pool_codes`` maps ``{pool_name: [loan_code, ...]}``. ``date_grid`` is the
    list of monthly ``pd.Timestamp`` the report's hist_bal_data uses. Returns a
    dict with ``warm_co``/``warm_rc`` (annual), ``warm_co_monthly``/
    ``warm_rc_monthly``, life-of-loan ``warm_co_totals``/``warm_rc_totals``/
    ``warm_net_co`` and per-pool ``hist_bal_data`` entries.
    """
    code2pool = {}
    for pool, codes in (pool_codes or {}).items():
        for c in codes:
            ci = _int_code(c)
            if ci is not None:
                code2pool[ci] = pool
    pools = list(pool_codes.keys())

    wb = openpyxl.load_workbook(warm_path, read_only=True, data_only=True)

    warm_co = defaultdict(lambda: defaultdict(float))
    warm_rc = defaultdict(lambda: defaultdict(float))
    co_mo = defaultdict(lambda: defaultdict(float))
    rc_mo = defaultdict(lambda: defaultdict(float))

    if 'CO Data' in wb.sheetnames:
        for r in wb['CO Data'].iter_rows(min_row=2, values_only=True):
            c = _int_code(r[2] if len(r) > 2 else None)
            if c not in code2pool:
                continue
            amt = r[3] if len(r) > 3 else None
            if not isinstance(amt, (int, float)):
                continue
            p = code2pool[c]
            y, ym = _year(r[4] if len(r) > 4 else None), _ym(r[4] if len(r) > 4 else None)
            if y:
                warm_co[y][p] += amt
            if ym:
                co_mo[ym][p] += amt

    if 'Recov Data' in wb.sheetnames:
        for r in wb['Recov Data'].iter_rows(min_row=2, values_only=True):
            c = _int_code(r[2] if len(r) > 2 else None)
            if c not in code2pool:
                continue
            amt = r[3] if len(r) > 3 else None
            if not isinstance(amt, (int, float)):
                continue
            p = code2pool[c]
            y, ym = _year(r[8] if len(r) > 8 else None), _ym(r[8] if len(r) > 8 else None)
            if y:
                warm_rc[y][p] += amt
            if ym:
                rc_mo[ym][p] += amt

    # Balances by quarter, grade-level, from the "<Mon>-<YY> Data" snapshots.
    qbal = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    for sn in wb.sheetnames:
        m = re.match(r'^([A-Z][a-z]{2})-(\d{2}) Data$', sn)
        if not m:
            continue
        yy, mm = 2000 + int(m.group(2)), _MON[m.group(1)]
        for r in wb[sn].iter_rows(min_row=2, values_only=True):
            c = _int_code(r[2] if len(r) > 2 else None)
            if c not in code2pool:
                continue
            bal = r[3] if len(r) > 3 else None
            if not isinstance(bal, (int, float)):
                continue
            g = _band(r[5] if len(r) > 5 else None, grades, no_score)
            qbal[code2pool[c]][(yy, mm)][g] += bal
    wb.close()

    # Snapshot-month grade balances from the current loan extract (df).
    snap_ts = pd.Timestamp(snapshot_date)
    snap_ym = (snap_ts.year, snap_ts.month)
    if df is not None and len(df):
        for pool in pools:
            pdf = df[df['loan_pool'] == pool]
            if not len(pdf):
                continue
            gb = defaultdict(float)
            for _, row in pdf.iterrows():
                gb[_band(row.get('current_fico_score'), grades, no_score)] += \
                    float(row.get('current_balance') or 0)
            if gb:
                qbal[pool][snap_ym] = gb

    grade_labels = [g.get('label') for g in grades] + [no_score]
    hist_bal = {}
    for pool in pools:
        qs = sorted(qbal[pool].keys())
        total, gseries = [], {gl: [] for gl in grade_labels}
        for d in date_grid:
            dym = (d.year, d.month)
            best = None
            for q in qs:
                if q <= dym:
                    best = q
                else:
                    break
            gb = qbal[pool].get(best, {}) if best else {}
            t = 0.0
            for gl in grade_labels:
                v = float(gb.get(gl, 0.0))
                gseries[gl].append(v)
                t += v
            total.append(t)
        gseries = {gl: v for gl, v in gseries.items() if any(v)}
        hist_bal[pool] = {'dates': list(date_grid), 'grades': gseries, 'total': total}

    # Life-of-loan windowed totals (monthly precision).
    warm_co_tot, warm_rc_tot, warm_net = {}, {}, {}
    for pool in pools:
        am = int((acl_months or {}).get(pool, 36) or 36)
        start = snap_ts - pd.DateOffset(months=am - 1)
        sy, sm = start.year, start.month

        def _win(mo):
            return sum(d.get(pool, 0.0) for (y, mth), d in mo.items()
                       if (y, mth) >= (sy, sm))
        cot, rct = _win(co_mo), _win(rc_mo)
        warm_co_tot[pool], warm_rc_tot[pool], warm_net[pool] = cot, rct, cot - rct

    return {
        'warm_co': {y: dict(d) for y, d in warm_co.items()},
        'warm_rc': {y: dict(d) for y, d in warm_rc.items()},
        'warm_co_monthly': {ym: dict(d) for ym, d in co_mo.items()},
        'warm_rc_monthly': {ym: dict(d) for ym, d in rc_mo.items()},
        'warm_co_totals': warm_co_tot,
        'warm_rc_totals': warm_rc_tot,
        'warm_net_co': warm_net,
        'hist_bal_data': hist_bal,
    }
