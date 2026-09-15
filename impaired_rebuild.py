"""Rebuild the impaired-loan calculation from the raw values the credit union
provides, applying the WARM workbook's "Impaired Loans" tab formulas in code.

The credit union sends only the raw inputs (Impairment Type, Member #, Suffix,
Loan Type, Current Balance, Days Delinquent, Balance at Other Lender, Collateral
Value, optional CU-provided amount, Notes -- columns A-J). The WARM's "Impaired
Loans" tab turns those into the Amount at Risk (LGD), Provision, and Balance
Removed figures that feed the report. This module reproduces that tab exactly so
the analyst never has to touch Excel.

Template = WARM "Impaired Loans" tab (source of truth; NOT the older standalone
"Specifically Identified Loans" template, whose provision percentages were stale).

Raw input columns (1-based):
    A Impairment Type   B Member #      C Loan Suffix   D Loan Type
    E Current Balance   F Days Delinq.  G Bal. at Other Lender (2nd mtg / HELOC)
    H Collateral Value (blank for unsecured)   I CU-provided amount   J Notes

Per-loan formulas replicated (WARM rows 35-333):
    K Member#-Suffix     = B & "-" & C
    L Total Loan         = E + G
    M LTV                = L / H            ("No Value" when H = 0)
    N LGD / Amt at Risk  = I if I given, else
                           0                when L-H <= 0
                           E                when L-H >= E
                           L-H              otherwise
    O Provision %        = 100% if I given, else provision-% table by type
    P Provision Amount   = N * O
    Q Balance Removed    = E               (the full current balance, always)

Provision-% table (WARM, from the tab's own A5:B9 -- Collateral Value calc only):
    Delinquent Loans 25%   Known Losses 100%   Repossessions 35%
    Foreclosed Real Estate 100%   Special Consideration 0.1%

Optionally resolves each loan's Loan Pool and Current Credit Grade against the
loan extract (as the WARM does via XLOOKUP into the 'Jun-26 Data' tab) to produce
the Balance-Removed-by-pool x grade pivot (the "Impaired Loans Pivot").

Usage:
    python impaired_rebuild.py "<path to CU's impaired .xlsx or the WARM workbook>"
"""

from __future__ import annotations

import os
from dataclasses import dataclass, field
from typing import Callable, Optional

import openpyxl


# The raw input columns (1-based). Same layout on the WARM "Impaired Loans" tab
# and the standalone "Spec Fund" file; only these are trusted, every calculated
# column is redone in code.
_COL = {
    "impairment_type": 1,   # A
    "member": 2,            # B
    "suffix": 3,            # C
    "loan_type": 4,         # D
    "current_balance": 5,   # E
    "days_delinquent": 6,   # F
    "balance_other": 7,     # G
    "collateral": 8,        # H
    "allowance_provided": 9,  # I
    "notes": 10,            # J
}

# Provision % by impairment type, per the WARM "Impaired Loans" tab (A5:B9).
# Used as the default when a file lacks a valid provision table of its own.
WARM_PROVISION_PCT = {
    "Delinquent Loans": 0.25,
    "Known Losses": 1.0,
    "Repossessions": 0.35,
    "Foreclosed Real Estate": 1.0,
    "Special Consideration": 0.001,
}

# Where the raw data rows live, per source tab.
_SHEET_RANGES = {
    "Impaired Loans": (35, 333),
    "Spec Fund": (31, 411),
}
_PROV_FIRST_ROW = 5     # provision-% table rows (A=type, B=percent)
_PROV_LAST_ROW = 19


def _num(v) -> float:
    """Coerce a cell to float; blanks/text -> 0.0."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return 0.0


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


@dataclass
class ImpairedLoan:
    impairment_type: str
    member: str
    suffix: str
    loan_type: str
    current_balance: float
    days_delinquent: float
    balance_other: float
    collateral: float
    allowance_provided: Optional[float]
    total_loans: float = 0.0
    ltv: Optional[float] = None
    amount_at_risk: float = 0.0
    percent: float = 0.0
    provision: float = 0.0
    balance_removed: float = 0.0
    pool: Optional[str] = None
    grade: Optional[str] = None

    @property
    def key(self) -> str:
        """Member#-Suffix, the WARM's XLOOKUP key (column K)."""
        return f"{self.member}-{self.suffix}"


@dataclass
class ImpairedResult:
    period: object = None
    provision_pct: dict = field(default_factory=dict)
    loans: list = field(default_factory=list)
    # per impairment type -> {'amount_at_risk', 'provision', 'balance_removed', 'count'}
    summary: dict = field(default_factory=dict)
    total_amount_at_risk: float = 0.0
    total_provision: float = 0.0
    total_balance_removed: float = 0.0
    # {pool: {grade: balance_removed}} once pool/grade are resolved
    pivot_balance_removed: dict = field(default_factory=dict)


def _norm(v) -> str:
    return str(v).strip().lower() if v is not None else ""


def _find_data_header(ws) -> Optional[int]:
    """Row index of the per-loan data header (Impairment Type + Member), or None.
    The raw data rows begin on the following row."""
    for r in range(1, min(ws.max_row, 80) + 1):
        if (_norm(ws.cell(row=r, column=1).value) == "impairment type"
                and _norm(ws.cell(row=r, column=2).value).startswith("member")):
            return r
    return None


def _read_provision_pct(ws) -> dict:
    """Locate the provision-% table (header 'Impairment Type | Provision Percentage')
    and read Type -> percent pairs below it. Falls back to the legacy fixed rows."""
    hdr = None
    for r in range(1, min(ws.max_row, 40) + 1):
        if (_norm(ws.cell(row=r, column=1).value) == "impairment type"
                and "provision" in _norm(ws.cell(row=r, column=2).value)):
            hdr = r
            break
    pct = {}
    rows = (range(hdr + 1, hdr + 30) if hdr is not None
            else range(_PROV_FIRST_ROW, _PROV_LAST_ROW + 1))
    for r in rows:
        name = ws.cell(row=r, column=1).value
        if _blank(name):
            if hdr is not None:
                break
            continue
        n = str(name).strip()
        if n.lower() == "total":
            break
        if n.upper() == "HIDE":
            continue
        pct[n] = _num(ws.cell(row=r, column=2).value)
    return pct


def _pick_sheet(wb):
    """Return (worksheet, (first_row, last_row)) for the impaired data-entry tab.

    Detects the tab robustly (its name may carry a leading space, an 'ASC 310-10'
    suffix, or be 'Spec Fund') by finding the one with a real data header, and
    derives the data range from that header -- rather than trusting fixed row
    positions. Helper tabs (Instructions / Management Adjustment / pivots) are
    skipped."""
    _skip = ("instruction", "management adjustment", "tdr", "pivot", "readme", "help")
    cands = []
    for name in wb.sheetnames:
        low = name.strip().lower()
        if any(s in low for s in _skip):
            continue
        ws = wb[name]
        hdr = _find_data_header(ws)
        if hdr is not None:
            score = (2 if ("impaired" in low or "spec fund" in low) else 1, ws.max_row)
            cands.append((score, name, ws, hdr))
    if cands:
        cands.sort(key=lambda x: x[0], reverse=True)
        _, _name, ws, hdr = cands[0]
        return ws, (hdr + 1, ws.max_row)
    for name in ("Impaired Loans", "Spec Fund"):
        if name in wb.sheetnames:
            return wb[name], _SHEET_RANGES[name]
    ws = wb[wb.sheetnames[0]]
    return ws, (2, ws.max_row)


def _read_period(ws):
    """Return the 'Report for Period Ending' value (scans for the label)."""
    for r in range(1, 12):
        if "report for period" in _norm(ws.cell(row=r, column=1).value):
            return ws.cell(row=r, column=2).value
    return ws.cell(row=3, column=2).value


def rebuild(
    path: str,
    provision_pct: Optional[dict] = None,
    resolve: Optional[Callable[[str, str, str], tuple]] = None,
) -> ImpairedResult:
    """Recompute the WARM "Impaired Loans" columns from the CU's raw inputs.

    ``provision_pct`` overrides the per-type percentages (defaults to the file's
    own table, falling back to the WARM values). ``resolve(member, suffix,
    loan_type) -> (pool, grade)`` optionally maps each loan to its pool/grade so
    the Balance-Removed pivot can be built.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, (first_row, last_row) = _pick_sheet(wb)

    prov = dict(provision_pct) if provision_pct else _read_provision_pct(ws)
    # If the file's table is missing/all-zero, fall back to the WARM percentages.
    if not prov or not any(v for v in prov.values()):
        prov = dict(WARM_PROVISION_PCT)
    res = ImpairedResult(provision_pct=prov)
    res.period = _read_period(ws)  # "Report for Period Ending"

    for r in range(first_row, last_row + 1):
        itype = ws.cell(row=r, column=_COL["impairment_type"]).value
        if _blank(itype):
            continue
        itype = str(itype).strip()
        if itype.upper() == "HIDE":
            continue
        E = _num(ws.cell(row=r, column=_COL["current_balance"]).value)
        G = _num(ws.cell(row=r, column=_COL["balance_other"]).value)
        H = _num(ws.cell(row=r, column=_COL["collateral"]).value)
        i_raw = ws.cell(row=r, column=_COL["allowance_provided"]).value
        overridden = not _blank(i_raw)

        L = E + G  # Total Loan (column L)
        # LGD / Amount at Risk (column N): CU override wins; else balance net of
        # collateral, floored at 0 and capped at the current balance.
        if overridden:
            N_at_risk = _num(i_raw)
        else:
            LH = L - H
            N_at_risk = 0.0 if LH <= 0 else (E if LH >= E else LH)
        # Provision % (column O): 100% when the CU provided the amount directly.
        O_pct = 1.0 if overridden else prov.get(itype, 0.0)
        P_prov = N_at_risk * O_pct
        member = str(ws.cell(row=r, column=_COL["member"]).value or "").strip()
        suffix = str(ws.cell(row=r, column=_COL["suffix"]).value or "").strip()
        loan_type = str(ws.cell(row=r, column=_COL["loan_type"]).value or "").strip()

        pool = grade = None
        if resolve is not None:
            try:
                pool, grade = resolve(member, suffix, loan_type)
            except Exception:
                pool = grade = None

        loan = ImpairedLoan(
            impairment_type=itype, member=member, suffix=suffix, loan_type=loan_type,
            current_balance=E,
            days_delinquent=_num(ws.cell(row=r, column=_COL["days_delinquent"]).value),
            balance_other=G, collateral=H,
            allowance_provided=None if not overridden else _num(i_raw),
            total_loans=L, ltv=(None if H == 0 else L / H),
            amount_at_risk=N_at_risk, percent=O_pct, provision=P_prov,
            balance_removed=E,  # WARM column Q = the full current balance, always
            pool=pool, grade=grade,
        )
        res.loans.append(loan)
        s = res.summary.setdefault(
            itype, {"amount_at_risk": 0.0, "provision": 0.0,
                    "balance_removed": 0.0, "count": 0})
        s["amount_at_risk"] += N_at_risk
        s["provision"] += P_prov
        s["balance_removed"] += E
        s["count"] += 1
        res.total_amount_at_risk += N_at_risk
        res.total_provision += P_prov
        res.total_balance_removed += E
        if pool is not None:
            g = grade if grade not in (None, "") else "Not Reported"
            res.pivot_balance_removed.setdefault(pool, {}).setdefault(g, 0.0)
            res.pivot_balance_removed[pool][g] += E
    return res


def _print(res: ImpairedResult) -> None:
    print(f"  Period: {res.period}")
    print(f"  Loans: {len(res.loans)}")
    hdr = (f"  {'Impairment Type':28}{'#':>5}{'LGD (Amt at Risk)':>20}"
           f"{'Prov %':>9}{'Provision':>15}{'Bal Removed':>16}")
    print(hdr)
    for itype, s in res.summary.items():
        pct = res.provision_pct.get(itype, 0.0)
        print(f"  {itype:28}{s['count']:>5}{s['amount_at_risk']:>20,.2f}"
              f"{pct:>9.2%}{s['provision']:>15,.2f}{s['balance_removed']:>16,.2f}")
    print(f"  {'TOTAL':28}{len(res.loans):>5}{res.total_amount_at_risk:>20,.2f}"
          f"{'':>9}{res.total_provision:>15,.2f}{res.total_balance_removed:>16,.2f}")
    if res.pivot_balance_removed:
        print("\n  Balance Removed by pool x grade:")
        for pool, grades in res.pivot_balance_removed.items():
            print(f"    {pool}: {sum(grades.values()):,.2f}")


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print('usage: python impaired_rebuild.py "<impaired .xlsx>"')
        raise SystemExit(2)
    _print(rebuild(sys.argv[1]))
